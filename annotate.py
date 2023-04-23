import pandas as pd
import os
import subprocess
import re
import cv2
import pytesseract
import configparser
import tkinter as tk
from tkinter import ttk
import time
from PIL import Image, ImageTk
import PIL.Image
import pickle
import shutil
import datetime
import sys
import openpyxl

# load config or create the file
# Set default values
config = configparser.ConfigParser()
config['Resource Paths'] = {
    'OCR_tesseract_path': 'C:/Program Files/Tesseract-OCR/tesseract.exe',
    'mpv_executable_path': 'C:/Program Files/MPV/mpv.exe',
    'video_folder_path': '',
    'annotation_file_path': ''
}
config['OCR settings'] = {
    'x_coordinate': '0',
    'y_coordinate': '0',
    'width': '500',
    'height': '40'
}
config['GUI settings'] = {
    'gui_dark_mode': '0',
    'mpv_width': '75%%',
    'mpv_height': '100%%',
    'mpv_x_coordinate': '-25',
    'mpv_y_coordinate': '0',
    'gui_width': '',
    'gui_height': '',
    'gui_x_coordinate': '0',
    'gui_y_coordinate': '0'
}
config['Workflow settings'] = {
    'Pause_time': '1',
    'Time_before_visit': '3',
    'Scan_default_folders': '1'
}

# Check if settings.ini exists, and create it with default values if not
if not os.path.exists('settings.ini'):
    with open('settings.ini', 'w', encoding='utf-8') as configfile:
        config.write(configfile)

# Read settings from settings.ini
config.read('settings.ini', encoding='utf-8')

# Get values from the config file
try:
    ocr_tesseract_path = config['Resource Paths'].get('OCR_tesseract_path', 'C:/Program Files/Tesseract-OCR/tesseract.exe').strip()
    mpv_executable_path = config['Resource Paths'].get('mpv_executable_path', 'C:/Program Files/MPV/mpv.exe').strip()
    video_folder_path = config['Resource Paths'].get('video_folder_path', '').strip()
    annotation_file_path = config['Resource Paths'].get('annotation_file_path', '').strip()
    scan_folders = config['Workflow settings'].get('Scan_default_folders', '0').strip()
except ValueError:
    print('Error: Invalid folder/file path found in settings.ini')
#Get gui values from config
try:
    mpv_width = config['GUI settings'].get('mpv_width', '75%%').strip()
    mpv_height = config['GUI settings'].get('mpv_height', '100%%').strip()
    mpv_x_coordinate = config['GUI settings'].get('mpv_x_coordinate', '0').strip()
    mpv_y_coordinate = config['GUI settings'].get('mpv_y_coordinate', '0').strip()
    pause_time = config['Workflow settings'].get('pause_time', '1').strip()
    time_before_visit = config['Workflow settings'].get('time_before_visit', '3').strip()
except ValueError:
    print('Error: Invalid window dimensions specified in settings.ini')
#It is sometimes necessary to specifiy the tesseracrt environmental variable in run-time
pytesseract.pytesseract.tesseract_cmd = ocr_tesseract_path

#Saving progress
MEMORY_FILE = 'memory.pkl'
# Check if the memory file exists
if os.path.isfile(MEMORY_FILE):
    # If it exists, load the data from the file
    with open(MEMORY_FILE, 'rb') as f:
        memory = pickle.load(f)
        m_index = memory.get('index', 0)  # Default to 0 if index key doesn't exist
        valid_annotations_array = memory.get('valid_annotations_array', [])
        # Ask the user if they want to continue where they left off
        answer = input(f"Do you want to continue from index {m_index}? (y/n) ")
        if answer.lower() == 'n':
            # If the user doesn't want to continue, overwrite the memory file with default values
            memory = {'index': 0, 'valid_annotations_array': []}
            with open(MEMORY_FILE, 'wb') as f:
                pickle.dump(memory, f)
            valid_annotations_array = []
            m_index = 0
        else:
            # If the user wants to continue, keep the current memory values
            pass
else:
    # If the memory file doesn't exist, create it with default values
    memory = {'index': 0, 'valid_annotations_array': []}
    with open(MEMORY_FILE, 'wb') as f:
        pickle.dump(memory, f)
    m_index = 0  # Set the index to the default value
    valid_annotations_array = []  # Set the list to the default value

#scan default folders
if scan_folders == "1":
    if not os.path.exists("videos/"):
        os.makedirs("videos/")
    if not os.path.exists("excel/"):
        os.makedirs("excel/")
    # Detect video files
    scan_video_files = [f for f in os.listdir('videos') if f.endswith('.mp4')]
    if scan_video_files:
        response = input(f"Video files detected in the default folder. Do you want to continue? (y/n): ")
        if response.lower() == 'y':
            video_folder_path = 'videos'
    # Detect Excel files
    scan_excel_files = [f for f in os.listdir('excel') if f.endswith('.xlsx') or f.endswith('.xls')]
    if scan_excel_files:
        response = input(f"Excel files detected in the default folder. Do you want to continue? (y/n): ")
        if response.lower() == 'y':
            print('Excel files in the folder:')
            for i, f in enumerate(scan_excel_files):
                print(f'{i + 1}. {f}')
            selection = int(input('Enter the number of the file you want to use: '))

            # Assign the path of the selected file to a variable
            if selection > 0 and selection <= len(scan_excel_files):
                annotation_file_path = os.path.join(("excel/"), scan_excel_files[selection - 1])
                print(f'Selected file: {annotation_file_path}')
            else:
                print('Invalid selection')

# set path to folder containing mp4 files
if not os.path.isdir(video_folder_path):
    video_folder_path = input("Enter the video folder path: ")
    if '"' in video_folder_path:
        video_folder_path = video_folder_path.replace('"', "")

# set path to excel file manually
if not os.path.isfile(annotation_file_path):
    annotation_file_path = input("Enter the path to the excel file: ")
    if '"' in annotation_file_path:
        annotation_file_path = annotation_file_path.replace('"',"")
#time_annotations = pd.read_csv(annotation_file_path, header=None, names=['time'], skiprows=1)

class RoundedButton(tk.Canvas):
    def __init__(self, master=None, text: str = "", radius=25, btnforeground="#000000", btnbackground="#ffffff",
                 clicked=None, *args, **kwargs):
        super(RoundedButton, self).__init__(master, *args, **kwargs)
        self.config(bg=self.master["bg"])
        self.btnbackground = btnbackground
        self.clicked = clicked

        self.radius = radius

        self.rect = self.round_rectangle(0, 0, 0, 0, tags="button", radius=radius, fill=btnbackground)
        self.text = self.create_text(0, 0, text=text, tags="button", fill=btnforeground, font=("Arial", 10),
                                     justify="center")

        self.tag_bind("button", "<ButtonPress>", self.border)
        self.tag_bind("button", "<ButtonRelease>", self.border)
        self.bind("<Configure>", self.resize)

        text_rect = self.bbox(self.text)
        if int(self["width"]) < text_rect[2] - text_rect[0]:
            self["width"] = (text_rect[2] - text_rect[0]) + 10

        if int(self["height"]) < text_rect[3] - text_rect[1]:
            self["height"] = (text_rect[3] - text_rect[1]) + 10
    def round_rectangle(self, x1, y1, x2, y2, radius=25, update=False,
                        **kwargs):  # if update is False a new rounded rectangle's id will be returned else updates existing rounded rect.
        points = [x1 + radius, y1,
                  x1 + radius, y1,
                  x2 - radius, y1,
                  x2 - radius, y1,
                  x2, y1,
                  x2, y1 + radius,
                  x2, y1 + radius,
                  x2, y2 - radius,
                  x2, y2 - radius,
                  x2, y2,
                  x2 - radius, y2,
                  x2 - radius, y2,
                  x1 + radius, y2,
                  x1 + radius, y2,
                  x1, y2,
                  x1, y2 - radius,
                  x1, y2 - radius,
                  x1, y1 + radius,
                  x1, y1 + radius,
                  x1, y1]
        if not update:
            return self.create_polygon(points, **kwargs, smooth=True)
        else:
            self.coords(self.rect, points)
    def resize(self, event):
        text_bbox = self.bbox(self.text)
        if self.radius > event.width or self.radius > event.height:
            radius = min((event.width, event.height))
        else:
            radius = self.radius
        width, height = event.width, event.height
        if event.width < text_bbox[2] - text_bbox[0]:
            width = text_bbox[2] - text_bbox[0] + 30
        if event.height < text_bbox[3] - text_bbox[1]:
            height = text_bbox[3] - text_bbox[1] + 30
        self.round_rectangle(5, 5, width - 5, height - 5, radius, update=True)
        bbox = self.bbox(self.rect)
        x = ((bbox[2] - bbox[0]) / 2) - ((text_bbox[2] - text_bbox[0]) / 2)
        y = ((bbox[3] - bbox[1]) / 2) - ((text_bbox[3] - text_bbox[1]) / 2)
        self.moveto(self.text, x, y)
    def border(self, event):
        if event.type == "4":
            self.itemconfig(self.rect, fill="#d2d6d3")
            if self.clicked is not None:
                self.clicked()
        else:
            self.itemconfig(self.rect, fill=self.btnbackground)
    def delete_button(self):
        self.delete(self.rect)
        self.delete(self.text)
        self.delete(self.tk)
def on_button_click(checkbox_vars, index):
    # Update the corresponding checkbox when a button is clicked
    checkbox_vars[index].set(1 - checkbox_vars[index].get())
def adjust_field(fields, index, value):
    current_value = int(fields[index]["text"])
    new_value = current_value + value
    fields[index].config(text=str(new_value))
def on_submit(root, checkbox_vars, comments_input, fields, buttons, close):
    # Get the values of the checkboxes, sliders, and input field and send them back to Python
    global input_data
    global close_script
    input_data = []
    input_data = [var.get() for var in checkbox_vars]
    input_data.append(comments_input.get())
    input_data.append(int(fields[0]["text"]))
    input_data.append(int(fields[1]["text"]))
    for i in range(len(buttons)):
        buttons[i].delete_button()
    root.destroy()
    if close:
        close_script = True
    else:
        close_script = False

def open_GUI(visit_data, index, video_path):

    #process data about visits
    this_visit_data = visit_data[index]
    minutes, seconds = divmod(this_visit_data[7], 60)
    this_time_string = f"{int(minutes):02d}:{int(seconds):02d}"
    this_timestamp = pd.to_datetime(this_visit_data[11], format='%Y%m%d_%H_%M_%S')
    if len(visit_data) > (index+1):
        next_visit_data = visit_data[index+1]
        minutes, seconds = divmod(next_visit_data[7], 60)
        next_time_string = f"{int(minutes):02d}:{int(seconds):02d}"
        next_timestamp = pd.to_datetime(next_visit_data[11], format='%Y%m%d_%H_%M_%S')
        print("WTF")
        # determine the delay before next visit
        if abs((next_timestamp - this_timestamp).total_seconds()) <= 5:
            text_field2_bg = 'orange'
        if abs((next_timestamp - this_timestamp).total_seconds()) <= 3:
            text_field2_bg = 'red'
        else:
            text_field2_bg = '#dddddd'
    else:
        next_visit_data = []
        for each in range(len(visit_data[index])):
            next_visit_data.append("NA")
        text_field2_bg = '#dddddd'
        next_time_string = "NA"
    #format based on dark_mode preference
    try:
        dark_mode = config['GUI settings'].get('gui_dark_mode', '0').strip()
    except ValueError:
        print('Error: Invalid GUI color mode specified in settings.ini')
    if dark_mode == "1":
        bg_col = '#3c3c3'
    else:
        bg_col = '#fdfdfd'

    #format gui window
    root = tk.Tk()
    root.title("Plant Visitor GUI")
    root.configure(background=bg_col)

    # Create a label for the text field
    text_label_frame = tk.Frame(root, background=root.cget('bg'))
    text_label_frame.pack(side=tk.TOP, padx=20, pady=0)
    text_label = tk.Label(text_label_frame, text="THIS VISIT", font=("Arial",15), background=root.cget('bg'))
    text_label.pack(side=tk.LEFT, padx=60)
    text_label2 = tk.Label(text_label_frame, text="NEXT VISIT", font=("Arial",15), background=root.cget('bg'))
    text_label2.pack(side=tk.RIGHT, padx=60)
    text_field_frame = tk.Frame(root, background=root.cget('bg'), highlightthickness=5, highlightbackground="#34495e")
    text_field_frame.pack(side=tk.TOP, padx=0, pady=0)

    # Create text fields to display visit info
    text_field = tk.Text(text_field_frame, height=5, width=33, font=("Arial",10))
    text_field.insert(tk.END, f"Time:\t\t{this_visit_data[3]}:{this_visit_data[4]}:{this_visit_data[5]}\nDuration:\t\t{this_time_string} \nVisitor:\t\t{this_visit_data[8]}\nFlowers visited:\t\t{this_visit_data[9]}")
    text_field.configure(state="disabled", highlightthickness=1, highlightbackground="#34495e", background="#39d2b4", relief="flat")
    text_field.pack(side=tk.LEFT, padx=(0,0))
    text_field2 = tk.Text(text_field_frame, height=5, width=33, font=("Arial",10))
    text_field2.insert(tk.END, f"Time:\t\t{next_visit_data[3]}:{next_visit_data[4]}:{next_visit_data[5]}\nDuration:\t\t{next_time_string} \nVisitor:\t\t{next_visit_data[8]}\nFlowers visited:\t\t{next_visit_data[9]}")
    text_field2.configure(highlightthickness=1, highlightbackground="#34495e", bg=text_field2_bg, relief="flat")
    text_field2.pack(side=tk.RIGHT, padx=(0,0))

    # Create a frame for the checkboxes and buttons
    checkbox_frame = tk.Frame(root)
    checkbox_frame.pack(side=tk.TOP, padx=20, pady=10)

    # Create 5 checkboxes and 5 buttons
    unchecked_image = PIL.Image.open("resources/img/no.png").resize((60, 40))
    unchecked_image = ImageTk.PhotoImage(unchecked_image)
    checked_image = PIL.Image.open("resources/img/yes.png").resize((60, 40))
    checked_image = ImageTk.PhotoImage(checked_image)
    #btn_image = PIL.Image.open("resources/img/btn.png").resize((300, 60))
    #btn_image = ImageTk.PhotoImage(btn_image)
    button_labels = ["fed on flower parts", "fed on pollen", "fed on nectar", "touched anthers", "touched stigmas"]
    checkbox_vars = [tk.IntVar() for _ in range(5)]
    checkboxes = []
    buttons = []
    fields = []
    for i in range(3):
        checkbox_frame = tk.Frame(root, background=root.cget('bg'))
        checkbox_frame.pack(side=tk.TOP, padx=20, pady=5)
        checkbox = tk.Checkbutton(checkbox_frame, image=unchecked_image, selectimage=checked_image, indicatoron=False, onvalue=1, offvalue=0, variable=checkbox_vars[i], background=root.cget('bg'), selectcolor=root.cget('bg'), bd=0)
        checkbox.config(highlightthickness=0, relief="flat")
        checkbox.pack(side=tk.LEFT)
        checkboxes.append(checkbox)
        #button = tk.Button(checkbox_frame, text=" ".join(["Visitor", button_labels[i]]),
                           #command=lambda i=i: on_button_click(i), image=btn_image, font=("Arial", 16), fg="#39d2b4", background=root.cget('bg'), activebackground=root.cget('bg'), relief="flat")
        button = RoundedButton(checkbox_frame, text=" ".join(["Visitor", button_labels[i]]), radius=90, btnbackground="#34495e", btnforeground="#39d2b4", clicked=lambda i=i: on_button_click(checkbox_vars, i), width=330, height=50)
        button.config(relief="flat", borderwidth=0, highlightthickness=0)
        button.pack(side=tk.RIGHT, padx=20)
        buttons.append(button)
    for i in range(2):
        #Create checkbox
        j = i + 3
        checkbox_frame = tk.Frame(root, background=root.cget('bg'))
        checkbox_frame.pack(side=tk.TOP, padx=20, pady=5)
        checkbox = tk.Checkbutton(checkbox_frame, image=unchecked_image, selectimage=checked_image, indicatoron=False, onvalue=1, offvalue=0, variable=checkbox_vars[j], background=root.cget('bg'), selectcolor=root.cget('bg'), bd=0, highlightthickness=0, relief="flat")
        checkbox.pack(side=tk.LEFT, padx=0)
        checkboxes.append(checkbox)
        # button = tk.Button(checkbox_frame, text=" ".join(["Visitor", button_labels[i]]),
        # command=lambda i=i: on_button_click(i), image=btn_image, font=("Arial", 16), fg="#39d2b4", background=root.cget('bg'), activebackground=root.cget('bg'), relief="flat")
        button = RoundedButton(checkbox_frame, text=" ".join(["Visitor", button_labels[j]]), radius=90, btnbackground="#34495e", btnforeground="#39d2b4", clicked=lambda j=j: on_button_click(checkbox_vars, j), width=220, height=50, relief="flat", borderwidth=0, highlightthickness=0)
        button.pack(side=tk.LEFT, padx=(30,0))
        buttons.append(button)
        button_minus = RoundedButton(checkbox_frame, text="-", radius=70,
                                     btnbackground="#34495e", btnforeground="#39d2b4",
                                     clicked=lambda i=i: adjust_field(fields, i, -1),
                                     width=50, height=50, relief="flat", borderwidth=0, highlightthickness=0)
        button_minus.pack(side=tk.LEFT, padx=0)
        buttons.append(button_minus)
        field = tk.Label(checkbox_frame, text="0", font=("Arial", 16), background=root.cget('bg'))
        field.pack(side=tk.LEFT, padx=0)
        fields.append(field)
        button_plus = RoundedButton(checkbox_frame, text="+", radius=70,
                                    btnbackground="#34495e", btnforeground="#39d2b4",
                                    clicked=lambda i=i: adjust_field(fields, i, +1),
                                    width=50, height=50, relief="flat", borderwidth=0, highlightthickness=0)
        button_plus.pack(side=tk.LEFT, padx=0)
        buttons.append(button_plus)

    # Create the comments section
    comments_frame = tk.Frame(root, background=root.cget('bg'))
    comments_frame.pack(side=tk.TOP, padx=20, pady=20)
    comments_label = tk.Label(comments_frame, text="Comments", font=("Arial", 10), background=root.cget('bg'))
    comments_label.pack(side=tk.LEFT)
    comments_input = tk.Entry(comments_frame, font=("Arial", 10), background=root.cget('bg'), width=50)
    comments_input.pack(side=tk.LEFT, padx=10, pady=10)

    # Create a submit button
    # Create a frame for the button
    submit_frame = tk.Frame(root, background=root.cget('bg'))
    submit_frame.pack(side=tk.TOP, padx=20, pady=0)
    submit_button = RoundedButton(submit_frame, text="Submit", clicked=lambda j=j: on_submit(root, checkbox_vars, comments_input, fields, buttons, False), radius=90, btnbackground="#34495e", btnforeground="#39d2b4", width=300, height=50, relief="flat", borderwidth=0, highlightthickness=0)
    submit_button.pack(side=tk.TOP, padx=20, pady=0)
    buttons.append(submit_button)
    sac_button = RoundedButton(submit_frame, text="Submit and Close",
                                  clicked=lambda j=j: on_submit(root, checkbox_vars, comments_input, fields, buttons, True),
                                  radius=90, btnbackground="#dddddd", btnforeground="#34495e", width=300, height=50,
                                  relief="flat", borderwidth=0, highlightthickness=0)
    sac_button.pack(side=tk.TOP, padx=20, pady=0)
    buttons.append(sac_button)


    #create a debugg console
    text_labeldc = tk.Label(root, text="DEBUG CONSOLE", font=("Arial", 15), background=root.cget('bg'))
    text_labeldc.pack(side=tk.TOP, padx=60, pady=5)
    text_labeldc_frame = tk.Frame(root, background=root.cget('bg'), highlightthickness=5, highlightbackground="#34495e")
    text_labeldc_frame.pack(side=tk.TOP, padx=20, pady=0)


    # Create text fields to display OCR info
    video_name = os.path.basename(video_path)[:-4]
    start_db_image = PIL.Image.open(f"OCR images/{video_name}_start.png").resize((450, 50))
    start_db_image = ImageTk.PhotoImage(start_db_image)
    text_labelsm = tk.Label(text_labeldc_frame, text="Start-time OCR input", font=("Arial", 12), background=root.cget('bg'))
    text_labelsm.pack(side=tk.TOP, padx=60, pady=0)
    start_image_label = tk.Label(text_labeldc_frame, image=start_db_image, font=("Arial", 15), background=root.cget('bg'))
    start_image_label.pack(side=tk.TOP, padx=(10,0))
    text_field3 = tk.Text(text_labeldc_frame, height=2, width=60, font=("Arial", 16))
    if (str(visit_data[index][13])[:2]) == "00":
        field_bg = "#ff0000"
    else:
        field_bg = "#34495e"
    text_field3.insert(tk.END,
                       f"Detected with OCR:\t{visit_data[index][13]}")
    text_field3.configure(state="disabled", highlightthickness=1, highlightbackground=field_bg,
                          background=root.cget('bg'),
                          relief="flat")
    text_field3.pack(side=tk.TOP, pady=5)

    end_db_image = PIL.Image.open(f"OCR images/{video_name}_end.png").resize((450, 50))
    end_db_image = ImageTk.PhotoImage(end_db_image)
    text_labelem = tk.Label(text_labeldc_frame, text="End-time OCR input", font=("Arial", 12),
                            background=root.cget('bg'))
    text_labelem.pack(side=tk.TOP, padx=60, pady=0)
    end_image_label = tk.Label(text_labeldc_frame, image=end_db_image, font=("Arial", 15), background='#ffffff')
    end_image_label.pack(side=tk.TOP, padx=(10,0))
    text_field4 = tk.Text(text_labeldc_frame, height=2, width=60, font=("Arial", 16))
    if (str(visit_data[index][14])[:2]) == "00":
        field_bg = "#ff0000"
    else:
        field_bg = "#34495e"
    text_field4.insert(tk.END,
                       f"Detected with OCR:\t{visit_data[index][14]}")
    text_field4.configure(state="disabled", highlightthickness=1, highlightbackground=field_bg,
                          background=root.cget('bg'), relief="flat")
    text_field4.pack(side=tk.TOP, pady=5)
    text_labelpr = tk.Label(text_labeldc_frame, text=f"Progress: {index+1}/{len(visit_data)}", font=("Arial", 10),
                            background=root.cget('bg'))
    text_labelpr.pack(side=tk.TOP, padx=60, pady=0)
    progress = ttk.Progressbar(text_labeldc_frame, orient=tk.HORIZONTAL, length=450, mode='determinate')
    progress.pack(padx=5, pady=5)
    progress["value"] =(100/len(visit_data))*(index+1)

    #open window
    try:
        gui_width = config['GUI settings'].get('gui_width', '').strip()
        gui_height = config['GUI settings'].get('gui_height', '').strip()
        gui_x_coordinate = config['GUI settings'].get('gui_x_coordinate', '').strip()
        gui_y_coordinate = config['GUI settings'].get('gui_y_coordinate', '').strip()
    except ValueError:
        print('Error: Invalid window dimensions specified in settings.ini')
    screen_width = root.winfo_screenwidth()
    screen_height = root.winfo_screenheight()
    if gui_width == '':
        if "%" in mpv_width:
            mpv_width_int = int(mpv_width.strip("%"))
            #gui_width = "".join([str(100 - mpv_width_int), "%"])
            gui_width = int((screen_width /(100/(100 - mpv_width_int))))
        else:
            mpv_width_int = int(mpv_width)
            gui_width = screen_width - mpv_width_int
    if gui_height == '':
        gui_height = screen_height
    if "%" in mpv_width:
        mpv_width_int = int(mpv_width.strip("%"))
        x = int((screen_width / (100/mpv_width_int)))
    else:
        x = int(screen_width - mpv_width_int)
    if not gui_x_coordinate == '':
        x = (x + int(gui_x_coordinate))
    y = 0
    if not gui_y_coordinate == '':
        y = 0 + int(gui_y_coordinate)
    root.geometry(f"{gui_width+55}x{gui_height}+{x-55}+{y}")
    root.lift()
    root.attributes('-topmost', True)
    root.after_idle(root.attributes, '-topmost', False)
    root.mainloop()
    return input_data

def get_text_from_video(video_filepath, start_or_end):
    # Read settings from settings.ini
    config.read('settings.ini', encoding='utf-8')
    try:
        x_coordinate = int(config['OCR settings'].get('x_coordinate', '0').strip())
        y_coordinate = int(config['OCR settings'].get('y_coordinate', '0').strip())
        width = int(config['OCR settings'].get('width', '500').strip())
        height = int(config['OCR settings'].get('height', '40').strip())
    except ValueError:
        # Handle cases where conversion to integer fails
        print('Error: Invalid integer value found in settings.ini')
    cap = cv2.VideoCapture(video_filepath)
    text_roi = (x_coordinate, y_coordinate, width, height)  # x, y, width, height
    if start_or_end == "end":
        total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
        second_to_last_frame_idx = total_frames - 5
        cap.set(cv2.CAP_PROP_POS_FRAMES, second_to_last_frame_idx)
    else:
        cap.set(cv2.CAP_PROP_POS_FRAMES, 24)
    ret, frame = cap.read()
    if ret:
        # Crop the image and pre-process it
        x, y, w, h = text_roi
        text_frame = frame[y:y + h, x:x + w]
        HSV_img = cv2.cvtColor(text_frame, cv2.COLOR_BGR2HSV)
        h, s, v = cv2.split(HSV_img)
        v = cv2.GaussianBlur(v, (1, 1), 0)
        thresh = cv2.threshold(v, 220, 255, cv2.THRESH_BINARY_INV)[1]
        kernel = cv2.getStructuringElement(cv2.MORPH_RECT, ksize=(1, 1))
        thresh = cv2.dilate(thresh, kernel)
        # text recognition
        OCR_text = pytesseract.image_to_string(thresh)
        # debug OCR file creation
        script_dir = os.path.dirname(os.path.abspath(__file__))
        output_dir = os.path.join(script_dir, "OCR images")
        if not os.path.exists(output_dir):
            os.makedirs(output_dir)
        OCR_file_name = ''.join([os.path.basename(video_filepath)[:-4], "_", start_or_end, ".png"])
        OCR_file_path = os.path.join(output_dir, OCR_file_name)
        cv2.imwrite(OCR_file_path, thresh)
    else:
        OCR_text = "none"
    cap.release()
    return OCR_text
def process_OCR_text(detected_text):
    if "\n" in detected_text:
        detected_text = detected_text.replace("\n", "")
    correct_format = r"\d{2}:\d{2}:\d{2}"
    if re.match(correct_format, detected_text[-8:]):
        print(' '.join(["Flow:", "Text detection successful -", detected_text[-8:]]))
        return_time = detected_text[-2:]
    else:
        print(' '.join(["Flow:", "Text detection failed -", detected_text]))
        print("Error: OCR detected text does not follow the expected format. The second value of this annotation shall set to arbitrary 0.")
        return_time = "00"
    return return_time
# define function to get start and end times for each video file
def get_video_start_end_times(video_filepath):
    video_filename = os.path.basename(video_filepath)
    print(' '.join(["Flow:", "Processing video file -" , video_filepath]))
    # get start time
    parts = video_filename[:-4].split("_")
    if len(parts) == 6:
        start_time_minutes = "_".join([parts[3], parts[4], parts[5]])
        print(' '.join(["Flow: Video name format with prefixes detected. Extracted the time values -", start_time_minutes]))
    else:
        print("Error: Some video file names have an unsupported format. Expected format is CO_LO1_SPPSPP1_YYYYMMDD_HH_MM. Script assumes format YYYYMMDD_HH_MM.")
        start_time_minutes = video_filename[:-4]
    #start_time_minutes = video_filename[:-4]
    text = get_text_from_video(video_filepath, "start")
    start_time_seconds = process_OCR_text(text)
    start_time_str = '_'.join([start_time_minutes, start_time_seconds])
    start_time = pd.to_datetime(start_time_str, format='%Y%m%d_%H_%M_%S')

    # get end time
    text = get_text_from_video(video_filepath, "end")
    end_time_seconds = process_OCR_text(text)
    end_time_str = pd.to_datetime('_'.join([start_time_minutes, end_time_seconds]), format='%Y%m%d_%H_%M_%S')
    end_time = end_time_str + pd.Timedelta(minutes=15)
    return start_time, end_time
def load_excel_table(file_path):
    # Define the file path
    # Define the columns to extract
    cols = [0, 1, 2, 3, 4, 5, 15, 18, 23, 27]
    # Read the Excel file, skipping the first two rows
    df = pd.read_excel(file_path, usecols=cols, skiprows=2, header=None)
    # Convert the month abbreviations in column 2 to month numbers
    months = {'Jan': '01', 'Feb': '02', 'Mar': '03', 'Apr': '04', 'May': '05', 'Jun': '06',
              'Jul': '07', 'Aug': '08', 'Sep': '09', 'Oct': '10', 'Nov': '11', 'Dec': '12'}
    print(df.iloc[:, 1])
    df.iloc[:, 1] = df.iloc[:, 1].replace(months)
    print(df.iloc[:, 1])
    df.iloc[:, 1] = df.iloc[:, 1].astype(int)
    # Add a column with the row number
    df['row_num'] = df.index + 1 + 2 # +2 is for skipped rows
    filtered_df = df[df.iloc[:, 6] == 1]
    filtered_df = filtered_df.copy()
    # print(filtered_df)
    # print(filtered_df.iloc[:, 5])
    for i in range(5):
        j = i + 1
        filtered_df.iloc[:, j] = filtered_df.iloc[:, j].astype(int).apply(lambda x: f'{x:02}')
    filtered_df.loc[:, 11] = filtered_df.iloc[:, 0:6].apply(lambda x: f"{x[0]}{x[1]}{x[2]}_{x[3]}_{x[4]}_{x[5]}",
                                                            axis=1)
    timestamps = filtered_df[11].values
    annotation_data_array = filtered_df.values.tolist()
    if not os.path.exists("resources/exc/"):
        # create directory
        os.makedirs("resources/exc/")
    filtered_df.to_excel("resources/exc/output_filtered.xlsx", index=False)
    return annotation_data_array

def append_excel_table(annotation_file_path, new_visit_data, annotation_data_array, index):
    # append excel file
    workbook = pd.read_excel(annotation_file_path, engine='openpyxl')
    #workbook = openpyxl.load_workbook(annotation_file_path)
    #worksheet = workbook.worksheets[0]
    cols = [35, 37, 41, 47, 49, 43, 48, 50]
    for j in range(len(new_visit_data)):
        workbook.at[annotation_data_array[index][10], cols[j]] = new_visit_data[j]
        #worksheet.cell(row=annotation_data_array[index][10], column=cols[j], value=new_visit_data[j])
        #workbook.save(annotation_file_path)
    workbook.to_excel(annotation_file_path, index=False, engine='openpyxl')

# Load videos
video_filepaths = []
for filename in os.listdir(video_folder_path):
    file_path = os.path.join(video_folder_path, filename)
    if os.path.isfile(file_path):
         video_filepaths.append(file_path)

# loop through time annotations and open corresponding video file
# extract video data beforehand to save processing time
if not len(valid_annotations_array) > 0:
    video_data = []
    valid_annotations_array = []
    valid_annotation_data_entry = []
    i: int
    for i, filepath in enumerate(video_filepaths):
        if filepath.endswith('.mp4'):
            video_start_time, video_end_time = get_video_start_end_times(video_filepaths[i])
            video_data_entry = [video_filepaths[i], video_start_time, video_end_time]
            video_data.append(video_data_entry)
    annotation_data_array = load_excel_table(annotation_file_path)
    for index, list in enumerate(annotation_data_array):
        print(' '.join(["Flow: Annotation number:", str(index+1)]))
        annotation_time = pd.to_datetime(annotation_data_array[index][11], format='%Y%m%d_%H_%M_%S')
        for i, list in enumerate(video_data):
            if video_data[i][1] <= annotation_time <= video_data[i][2]:
                for each in range(len(annotation_data_array[index])):
                    valid_annotation_data_entry.append(annotation_data_array[index][each])
                for each in range(3):
                    valid_annotation_data_entry.append(video_data[i][each])
                print(valid_annotation_data_entry)
                valid_annotations_array.append(valid_annotation_data_entry)
                valid_annotation_data_entry = []
    memory['valid_annotations_array'] = valid_annotations_array
    with open(MEMORY_FILE, 'wb') as f:
        pickle.dump(memory, f)
    print(valid_annotations_array)
for index in range(m_index, len(valid_annotations_array)):
    print(f"Processing item {index}")
    annotation_time = pd.to_datetime(valid_annotations_array[index][11], format='%Y%m%d_%H_%M_%S')
    if valid_annotations_array[index][13] <= (annotation_time - pd.Timedelta(seconds=int(time_before_visit))):
        annotation_offset = ((annotation_time - pd.Timedelta(seconds=int(time_before_visit))) - valid_annotations_array[index][13]).total_seconds()
    else:
        annotation_offset = (annotation_time - valid_annotations_array[index][13]).total_seconds()
    mpv_process = subprocess.Popen([mpv_executable_path, f'--geometry={mpv_width}x{mpv_height}+{mpv_x_coordinate}+{mpv_y_coordinate}', '-ss', str(annotation_offset), valid_annotations_array[index][12]])
    new_visit_data = open_GUI(valid_annotations_array, index, valid_annotations_array[index][12])
    append_excel_table(annotation_file_path, new_visit_data, valid_annotations_array, index)
    mpv_process.terminate()
    time.sleep(int(pause_time))
    if len(valid_annotations_array) >= (index + 1):
        memory['index'] = index+1
        with open(MEMORY_FILE, 'wb') as f:
            pickle.dump(memory, f)
    else:
        break
    if close_script:
        sys.exit()
if os.path.isfile(MEMORY_FILE):
    if not os.path.isdir("resources/memory"):
        os.makedirs("resources/memory")
    now = datetime.datetime.now().strftime('%Y-%m-%d_%H-%M-%S')
    new_filename = f'memory_{now}.pkl'
    shutil.copy2(MEMORY_FILE, os.path.join("resources/memory", new_filename))
    os.remove(MEMORY_FILE)
    print("Task completed. Memory file deleted. Backup can be found in resources/memory folder.")
